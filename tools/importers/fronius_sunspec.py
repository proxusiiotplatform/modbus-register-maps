#!/usr/bin/env python3
"""Convert a Fronius SunSpec register map workbook (the "Complete Map" sheet) into a device map.

    python3 tools/importers/fronius_sunspec.py <Register_Map.xlsx> <header.yaml> <out.yaml>

Fronius numbers registers from 1: the SunSpec marker 'SunS' that the SunSpec reference client
(pysunspec2) reads at address 40000 is listed at 40001. So address = Start - 1.
Each SunSpec model becomes a group named after the workbook's per-model sheet; registers keep
their SunSpec point name in the notes. Scale-factor points (sunssf) are linked with scale_factor.

Needs openpyxl==3.1.5.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

TYPES = {
    "uint16": "uint16", "int16": "int16", "uint32": "uint32", "int32": "int32", "float32": "float32",
    "enum16": "uint16", "count": "uint16", "sunssf": "int16", "acc32": "uint32", "acc64": "uint64",
    "bitfield16": "bitfield16", "bitfield32": "bitfield", "string": "string", "uint64": "uint64",
}
ACCESS = {"R": "r", "RW": "rw", "W": "w"}


def text(value) -> str:
    return "" if value is None else str(value).strip()


def snake(value: str) -> str:
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", value)
    value = re.sub(r"[^a-zA-Z0-9]+", "_", value).strip("_").lower()
    return ("p_" + value if value[:1].isdigit() else value)[:64].rstrip("_")


def q(value: str) -> str:
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def convert(source: Path) -> tuple[list[dict], dict]:
    book = openpyxl.load_workbook(str(source), read_only=True, data_only=True)
    models = [re.sub(r"\s*\(.*\)$", "", name) for name in book.sheetnames
              if name not in ("Complete Map", "End Block")]
    rows = list(book["Complete Map"].iter_rows(values_only=True))
    registers, keys = [], set()
    stats = {"pad": 0, "end_block": 0}
    model_index, group = -1, "header"
    for row in rows:
        start = text(row[0])
        if not start.isdigit():
            continue
        start, size = int(start), int(text(row[2]) or 1)
        access_raw, point, desc, typ_raw = text(row[3]), text(row[5]), text(row[6]), text(row[7])
        units, sf = text(row[8]), text(row[9])
        if point == "ID" and start > 40001:
            model_index += 1
            group = snake(models[model_index]) if model_index < len(models) else "end"
        if group == "end":
            stats["end_block"] += 1
            continue
        if typ_raw == "pad":
            stats["pad"] += 1
            continue
        typ = TYPES[typ_raw]
        key = snake(f"{group} {point}") if group != "header" else snake(point)
        if key in keys:
            key = f"{key}_{start - 1}"
        keys.add(key)
        name = desc if desc and len(desc) <= 60 and "\n" not in desc else point
        item: dict = {"key": key, "name": name, "group": group, "table": "holding", "address": start - 1, "type": typ}
        if typ in ("string", "bitfield"):
            item["length"] = size
        if sf and sf != "None":
            item["scale_factor"] = snake(f"{group} {sf}")
        if units and units != "None":
            item["unit"] = units
        item["access"] = ACCESS.get(access_raw, "r")
        notes = [f"SunSpec point {point}."]
        if typ_raw == "sunssf":
            notes.append("Scale factor: a power-of-ten exponent for the registers that name it.")
        elif typ_raw in ("acc32", "acc64"):
            notes.append("Accumulator: counts up and wraps.")
        elif typ_raw == "enum16":
            notes.append("Enumerated value; see the Fronius Modbus TCP & RTU manual for the states.")
        item["notes"] = " ".join(notes)
        registers.append(item)
    return registers, stats


def emit(item: dict) -> str:
    parts = []
    for k, v in item.items():
        if isinstance(v, int):
            parts.append(f"{k}: {v}")
        else:
            parts.append(f"{k}: {q(v) if k in ('name', 'unit', 'notes') else v}")
    return "  - {" + ", ".join(parts) + "}"


def main(argv: list[str]) -> int:
    if len(argv) != 4:
        print(__doc__)
        return 2
    source, header, out = map(Path, argv[1:])
    registers, stats = convert(source)
    head = header.read_text(encoding="utf-8").rstrip("\n")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(f"{head}\nregisters:\n" + "\n".join(emit(r) for r in registers) + "\n", encoding="utf-8")
    print(f"{out}: {len(registers)} registers, {stats}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
